from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from .forms import FornecedorForm
from .models import Fornecedor
from django.contrib import messages
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from django.http import HttpResponse

@login_required
def CadFornecedor(request):
    fornecedores = Fornecedor.objects.all()
    
    # Filtro por nome do fornecedor
    nome = request.GET.get('fornecedor')
    if nome:
        fornecedores = fornecedores.filter(fornecedor__icontains=nome)
    
    # Filtro por email
    email = request.GET.get('email')
    if email:
        fornecedores = fornecedores.filter(fornecedor_email__icontains=email)
        
    # Filtro por telefone
    telefone = request.GET.get('telefone')
    if telefone:
        fornecedores = fornecedores.filter(fornecedor_fone__icontains=telefone)
        
    # Exportar Excel se solicitado
    if request.GET.get('exportar_excel'):
        wb = gerar_excel_fornecedores(fornecedores)
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="fornecedores.xlsx"'
        wb.save(response)
        return response

    return render(request, 'CadFornecedor.html', {
        'fornecedores': fornecedores
    })

@login_required
def cria_fornecedor(request):
    if request.method == 'POST':
        form = FornecedorForm(request.POST)
        if form.is_valid():
            fornecedor = form.save(commit=False)
            fornecedor.save()
            return redirect('CadFornecedor')
    else:
        # print('Faltando dados:')
        form = FornecedorForm()

    return render(request, 'cria_fornecedor.html', {'form': form})


@login_required  # Add login requirement for security
def excluir_fornecedor(request, id):
    fornecedor = get_object_or_404(Fornecedor, id=id)
    
    if request.method == 'POST':
        fornecedor.delete()
        return redirect('CadFornecedor')
    
    # Create a separate template for deletion confirmation
    return render(request, 'confirmar_exclusao.html', {
        'fornecedor': fornecedor
    })

@login_required
def editar_fornecedor(request, id):
    fornecedor = get_object_or_404(Fornecedor, id=id)

    if request.method == 'POST':
        form = FornecedorForm(request.POST, instance=fornecedor)
        if form.is_valid():
            form.save()
            return redirect('CadFornecedor')
    else:
        form = FornecedorForm(instance=fornecedor)

    return render(request, 'editar_fornecedor.html', {'form': form})

def gerar_excel_fornecedores(fornecedores):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fornecedores"

    headers = ["Nome do Fornecedor", "Telefone", "Site", "Email"]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="007BFF", end_color="007BFF", fill_type="solid")
    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    alignment = Alignment(horizontal='center', vertical='center')

    # Add header row
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = alignment

    row_num = 2
    for f in fornecedores:
        ws.cell(row=row_num, column=1, value=f.fornecedor)
        ws.cell(row=row_num, column=2, value=f.fornecedor_fone)
        ws.cell(row=row_num, column=3, value=f.fornecedor_site)
        ws.cell(row=row_num, column=4, value=f.fornecedor_email)
        row_num += 1

    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max_length + 2

    return wb