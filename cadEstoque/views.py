from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from .forms import PecasForm
from .models import Pecas
from CadFornecedor.models import Fornecedor
from django.db.models import Q
import logging
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from django.http import HttpResponse

logger = logging.getLogger(__name__)

def gerar_excel_pecas(pecas):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Estoque de Peças"
    headers = ["Cadastro Nº", "Descrição", "Código", "Fornecedor", "Estoque", "Preço"]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="007BFF", end_color="007BFF", fill_type="solid")
    border = Border(left=Side(style='thin'), right=Side(style='thin'),
                    top=Side(style='thin'), bottom=Side(style='thin'))
    alignment = Alignment(horizontal='center', vertical='center')
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = alignment

    row_number = 2
    for peca in pecas:
        ws.cell(row=row_number, column=1, value=peca.id)
        ws.cell(row=row_number, column=2, value=peca.descricao)
        ws.cell(row=row_number, column=3, value=peca.codigo)
        ws.cell(row=row_number, column=4, value=str(peca.fornecedor))
        ws.cell(row=row_number, column=5, value=peca.estoque_atual)
        ws.cell(row=row_number, column=6, value=peca.preco)
        row_number += 1

    # Optionally adjust column widths...
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                if cell.value and len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        ws.column_dimensions[col_letter].width = max_length + 2

    return wb

@login_required
def teste(request):
    pecas = Pecas.objects.all()
    fornecedor_id = request.POST.get('fornecedor')
    # print("Request POST:", request.POST)  # Debug do formulário enviado
    # print("Fornecedor ID recebido:", fornecedor_id)  # Debug do fornecedor

    # Recupera os filtros da requisição GET
    descricao = request.GET.get('descricao')
    codigo = request.GET.get('codigo')

    # Armazena filtros na sessão
    if descricao or codigo:
        request.session['descricao'] = descricao
        request.session['codigo'] = codigo
        return redirect('cadEstoque')  # Redireciona para limpar os campos da URL

    # Recupera os filtros da sessão, se existentes
    descricao = request.session.get('descricao')
    codigo = request.session.get('codigo')

    # Aplica os filtros
    if descricao:
        pecas = pecas.filter(descricao__icontains=descricao)
    if codigo:
        pecas = pecas.filter(codigo__icontains=codigo)

    # Adiciona URL da foto para cada peça
    for peca in pecas:
        peca.foto_url = peca.foto.url if peca.foto else None

    # Limpa os filtros da sessão para novas buscas
    request.session.pop('descricao', None)
    request.session.pop('codigo', None)

    # If exporting excel then generate the file
    if request.GET.get('exportar_excel'):
        wb = gerar_excel_pecas(pecas)
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="pecas.xlsx"'
        wb.save(response)
        return response

    return render(request, 'teste.html', {'pecas': pecas})


@login_required
def cria_peca(request):
    if request.method == 'POST':
        logger.debug(f"Dados recebidos no POST: {request.POST}")
        form = PecasForm(request.POST, request.FILES)
        fornecedor_id = request.POST.get('fornecedor')
        # print("Request POST:", request.POST)  # Debug do formulário enviado
        # print("Fornecedor ID recebido:", fornecedor_id)  # Debug do fornecedor
        if form.is_valid():
            # Verifica fornecedor
            peca = form.save(commit=False)
            if fornecedor_id:
                try:
                    peca.fornecedor = Fornecedor.objects.get(id=fornecedor_id)
                except Fornecedor.DoesNotExist:
                    form.add_error('fornecedor', 'Fornecedor selecionado não existe.')

            logger.debug("Formulário validado com sucesso.")
            fornecedor_id = request.POST.get('fornecedor')
            # print("Fornecedor ID antes de salvar:", fornecedor_id)
            peca = form.save()
            logger.debug(f"Peça criada: {peca}")
            return redirect('cadEstoque')
        else:
            logger.error(f"Erros no formulário: {form.errors}")
    else:
        form = PecasForm()

    todos_fornecedores = Fornecedor.objects.all()
    return render(request, 'cria_peca.html', {'form': form, 'fornecedores': todos_fornecedores})


@login_required
def editar_peca(request, id):
    peca = get_object_or_404(Pecas, id=id)
    fornecedor_selecionado = peca.fornecedor  # Obtem o fornecedor diretamente, se existir
    # Adicionando prints para debug
    # print("\n=== DEBUG INFORMAÇÕES DE FORNECEDORES ===")
    # print(f"Fornecedor selecionado: {fornecedor_selecionado}")
    # if fornecedor_selecionado:
        # print(f"Detalhes do fornecedor selecionado:")
        # print(f"- ID: {fornecedor_selecionado.id}")
        # print(f"- Nome: {fornecedor_selecionado.fornecedor}")
        # print(f"- Email: {fornecedor_selecionado.fornecedor_email}")
        # print(f"- Telefone: {fornecedor_selecionado.fornecedor_fone}")
    if request.method == 'POST':
        form = PecasForm(request.POST, request.FILES, instance=peca)
        fornecedor_id = request.POST.get('fornecedor')
        if form.is_valid():
            peca = form.save(commit=False)
            if fornecedor_id:
                try:
                    fornecedor = Fornecedor.objects.get(id=fornecedor_id)
                    peca.fornecedor = fornecedor
                except Fornecedor.DoesNotExist:
                    form.add_error('fornecedor', 'Fornecedor selecionado não existe.')
            form.save()
            return redirect('cadEstoque')
    else:
        form = PecasForm(instance=peca)

    todos_fornecedores = Fornecedor.objects.all()
    return render(request, 'edita_peca.html', {'form': form, 'fornecedores': todos_fornecedores,
        'fornecedor_selecionado': fornecedor_selecionado})

@login_required
def excluir_peca(request, id):
    peca = get_object_or_404(Pecas, id=id)
    if request.method == 'POST':
        peca.delete()
        return redirect('cadEstoque')
    return render(request, 'confirmar_exclusao.html', {'peca': peca})

@login_required
def busca_fornecedores(request):
    nome_parcial = request.GET.get('nome', '')
    fornecedores = Fornecedor.objects.filter(
        Q(fornecedor__icontains=nome_parcial)
    )
    return render(request, 'busca_fornecedores.html', {'fornecedores': fornecedores})

